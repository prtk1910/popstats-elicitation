"""Minimal Markdown -> DOCX converter (headings, paragraphs, bullets, tables, images)."""

from __future__ import annotations

import re
from pathlib import Path


def md_to_docx(md_path: Path, docx_path: Path) -> Path:
    from docx import Document
    from docx.shared import Inches

    doc = Document()
    lines = md_path.read_text().splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        img = re.fullmatch(r"!\[.*?\]\((.+?)\)", stripped)
        if img:
            src = md_path.parent / img.group(1)
            png = src.with_suffix(".png")
            target = png if png.exists() else src
            if target.exists() and target.suffix.lower() == ".png":
                try:
                    doc.add_picture(str(target), width=Inches(5.8))
                except Exception:
                    pass
            i += 1
            continue

        m = re.match(r"^(#{1,4})\s+(.*)", stripped)
        if m:
            doc.add_heading(m.group(2), level=len(m.group(1)))
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and \
                re.fullmatch(r"\|[\s:\-|]+\|", lines[i + 1].strip()):
            header = [c.strip() for c in stripped.strip("|").split("|")]
            rows = []
            j = i + 2
            while j < len(lines) and lines[j].strip().startswith("|"):
                rows.append([c.strip() for c in lines[j].strip().strip("|").split("|")])
                j += 1
            table = doc.add_table(rows=1 + len(rows), cols=len(header))
            table.style = "Light Grid Accent 1"
            for c, text in enumerate(header):
                table.rows[0].cells[c].text = text.replace("**", "")
            for r, row in enumerate(rows):
                for c, text in enumerate(row):
                    if c < len(table.rows[r + 1].cells):
                        table.rows[r + 1].cells[c].text = text.replace("**", "")
            i = j
            continue

        if stripped.startswith("- "):
            p = doc.add_paragraph(stripped[2:], style="List Bullet")
            i += 1
            continue

        doc.add_paragraph(stripped)
        i += 1

    doc.save(docx_path)
    return docx_path
